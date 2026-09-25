# -*- coding: utf-8 -*-
"""Génère le document Word de présentation e-Justice / e-Detenu."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE = RGBColor(0x00, 0x4A, 0x94)
BLUE_DEEP = RGBColor(0x00, 0x33, 0x66)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
RED = RGBColor(0xCE, 0x10, 0x20)
INK = RGBColor(0x12, 0x26, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="004a94", sz="8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, size=10, bold=False, color=INK, align="left", fill=None):
    if fill:
        shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_borders(cell, "d5dee8", "4")


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "0066C7")
        pBdr.append(bottom)
        pPr.append(pBdr)
        run = p.add_run(text)
        set_run_font(run, size=16, bold=True, color=BLUE)
    else:
        run = p.add_run(text)
        set_run_font(run, size=13, bold=True, color=BLUE_DEEP)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead:
        r1 = p.add_run("•  " + bold_lead)
        set_run_font(r1, size=11, bold=True, color=BLUE)
        r2 = p.add_run(text)
        set_run_font(r2, size=11, color=INK)
    else:
        r = p.add_run("•  " + text)
        set_run_font(r, size=11, color=INK)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=MUTED)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=10, bold=True, color=WHITE, align="center", fill="004A94")
    for r_i, row in enumerate(rows):
        fill = "F4F8FC" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, size=10, color=INK, fill=fill)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    return table


def footer_setup(section, text):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text + "   ·   ")
    set_run_font(run, size=8, color=MUTED)
    # PAGE field
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r_page = p.add_run()
    set_run_font(r_page, size=8, color=MUTED)
    r_page._r.append(fld1)
    r_page._r.append(instr)
    r_page._r.append(fld2)
    run2 = p.add_run(" / ")
    set_run_font(run2, size=8, color=MUTED)
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    r_np = p.add_run()
    set_run_font(r_np, size=8, color=MUTED)
    r_np._r.append(fld3)
    r_np._r.append(instr2)
    r_np._r.append(fld4)


def banner_table(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "004A94")
    cell.text = ""
    for i, (txt, size, bold) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6 if i == 0 else 2)
        p.paragraph_format.space_after = Pt(6 if i == len(lines) - 1 else 2)
        run = p.add_run(txt)
        set_run_font(run, size=size, bold=bold, color=WHITE)
    set_cell_borders(cell, "F7D116", "16")


def gold_line(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "F7D116")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(2.0)
    footer_setup(section, "e-Justice  ·  Module e-Detenu  ·  République démocratique du Congo")

    core = doc.core_properties
    core.title = "Présentation de la plateforme e-Justice et du module e-Detenu"
    core.author = "e-Justice — République démocratique du Congo"
    core.subject = "Document de présentation institutionnelle"
    core.language = "fr-FR"

    # ----- Couverture -----
    for _ in range(3):
        doc.add_paragraph()
    banner_table(doc, [
        ("RÉPUBLIQUE DÉMOCRATIQUE DU CONGO", 11, True),
        ("Justice  ·  Administration pénitentiaire", 10, False),
    ])
    gold_line(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("e-Justice")
    set_run_font(r, size=36, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Module e-Detenu")
    set_run_font(r, size=22, bold=True, color=BLUE_DEEP)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Système national de gestion des centres pénitenciers")
    set_run_font(r, size=13, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("DOCUMENT DE PRÉSENTATION")
    set_run_font(r, size=12, bold=True, color=RED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("Septembre 2026   ·   Version 1.0")
    set_run_font(r, size=11, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Usage institutionnel")
    set_run_font(r, size=10, color=MUTED)

    doc.add_page_break()

    # ----- Sommaire -----
    add_heading_custom(doc, "Sommaire", 1)
    sommaire = [
        "1.  Objet du document",
        "2.  Contexte et enjeux",
        "3.  La plateforme e-Justice",
        "4.  Le module e-Detenu",
        "5.  Gouvernance et acteurs",
        "6.  Chaîne opérationnelle",
        "7.  Fonctionnalités du module",
        "8.  Couverture territoriale",
        "9.  Architecture, sécurité et identité visuelle",
        "10. État d’avancement et perspectives",
        "11. Conclusion",
    ]
    for item in sommaire:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(item)
        set_run_font(r, size=12, color=BLUE_DEEP)

    doc.add_page_break()

    # ----- 1 -----
    add_heading_custom(doc, "1.  Objet du document", 1)
    add_body(
        doc,
        "Le présent document décrit, dans un ordre logique, la plateforme nationale e-Justice "
        "et son premier module opérationnel, e-Detenu. Il s’adresse aux autorités de tutelle, "
        "aux administrateurs de centres pénitenciers et aux équipes chargées du déploiement.",
    )
    add_body(
        doc,
        "Il présente d’abord le cadre institutionnel, puis la vision d’ensemble de e-Justice, "
        "ensuite le périmètre d’e-Detenu, la répartition des responsabilités, le déroulement "
        "métier, les fonctions disponibles, la couverture des établissements, les choix "
        "techniques et, enfin, les perspectives d’extension.",
    )

    # ----- 2 -----
    add_heading_custom(doc, "2.  Contexte et enjeux", 1)
    add_body(
        doc,
        "La République démocratique du Congo dispose d’un réseau d’établissements pénitentiaires "
        "répartis sur l’ensemble du territoire. La gestion quotidienne — identification des "
        "détenus, occupation des cellules, visites, personnel, surveillance — reste souvent "
        "fragmentée entre supports papier et outils locaux non interconnectés.",
    )
    add_body(doc, "e-Detenu répond à quatre exigences :")
    add_bullet(doc, " disposer d’une fiche unique, à jour, pour chaque personne détenue ;", "Fiabiliser :")
    add_bullet(doc, " séparer clairement le pilotage national de la gestion de chaque prison ;", "Séparer :")
    add_bullet(doc, " connaître l’occupation réelle des centres et les mouvements (incarcération, libération, transfert) ;", "Piloter :")
    add_bullet(doc, " préparer l’intégration future avec le parquet, le greffe et le casier judiciaire.", "Inscrire :")

    # ----- 3 -----
    add_heading_custom(doc, "3.  La plateforme e-Justice", 1)
    add_body(
        doc,
        "e-Justice est conçue comme la porte d’entrée numérique de la chaîne pénale et civile. "
        "e-Detenu en constitue le maillon « exécution des peines ». Les autres fonctions sont "
        "affichées sur la page d’accueil afin de montrer la trajectoire complète, même lorsqu’elles "
        "ne sont pas encore déployées.",
    )
    add_heading_custom(doc, "3.1  Les dix fonctions prévues", 2)
    add_table(
        doc,
        ["Fonction", "Domaine", "Statut"],
        [
            ["Police judiciaire", "Enquêtes, PV, scellés, interpellations", "Prévu"],
            ["Poursuites publiques", "Parquet, classements, réquisitoires", "Prévu"],
            ["Instruction préparatoire", "Mandats, commissions rogatoires", "Prévu"],
            ["Audiences et greffe", "Rôle, convocations, minutes", "Prévu"],
            ["Jugements et arrêts", "Rédaction, signification, recours", "Prévu"],
            ["Casier judiciaire", "Mentions, extraits, réhabilitation", "Prévu"],
            ["e-Detenu", "Centres, fiches, visites, personnel, surveillance", "Opérationnel"],
            ["Justice civile", "Assignations, référés, contentieux", "Prévu"],
            ["Aide juridictionnelle", "Demandes d’aide, commission d’office", "Prévu"],
            ["Pilotage", "Statistiques nationales et délais de justice", "Prévu"],
        ],
        [4.2, 9.5, 3.3],
    )
    add_caption(doc, "Tableau 1 — Architecture fonctionnelle de e-Justice")
    add_body(
        doc,
        "Cette organisation permet d’éviter les silos : lorsqu’un dossier judiciaire aboutit "
        "à une incarcération, e-Detenu en assure l’exécution ; à terme, le casier et les "
        "statistiques nationales s’alimenteront des mêmes données.",
    )

    # ----- 4 -----
    add_heading_custom(doc, "4.  Le module e-Detenu", 1)
    add_body(
        doc,
        "e-Detenu est le système de gestion des centres pénitenciers. Il couvre l’identité "
        "du détenu, le dossier judiciaire joint à l’écrou, l’affectation en cellule selon "
        "la structure de la prison, les visites, le personnel du centre, la surveillance "
        "et les documents officiels (fiche d’identité, listes PDF).",
    )
    add_heading_custom(doc, "4.1  Principes directeurs", 2)
    add_bullet(doc, " un administrateur central voit l’ensemble du réseau, crée les prisons et nomme les administrateurs de centre ; il ne saisit pas le quotidien d’une prison.", "National :")
    add_bullet(doc, " chaque admin de prison n’agit que sur son établissement (détenus, personnel, visites).", "Local :")
    add_bullet(doc, " les écrans sont conçus pour un usage de guichet : listes compactes, fiches à onglets, actions immédiates.", "Opérationnel :")
    add_bullet(doc, " l’interface reprend les couleurs du drapeau national (bleu, jaune, rouge) pour un usage institutionnel.", "Souverain :")

    # ----- 5 -----
    add_heading_custom(doc, "5.  Gouvernance et acteurs", 1)
    add_body(
        doc,
        "Les droits d’accès suivent le rôle métier. Un même logiciel sert le niveau national "
        "et le niveau de la prison, sans mélanger les écritures.",
    )
    add_table(
        doc,
        ["Acteur", "Rôle dans e-Detenu", "Peut créer / modifier"],
        [
            [
                "Administrateur central",
                "Pilotage national, consultation de toutes les prisons",
                "Centres pénitenciers et administrateurs de prison",
            ],
            [
                "Administrateur de prison",
                "Direction numérique de son centre",
                "Détenus, personnel, visites et visiteurs du centre",
            ],
            [
                "Directeur",
                "Encadrement du centre auquel il est rattaché",
                "Données opérationnelles de son centre",
            ],
            [
                "Agent",
                "Saisie quotidienne (écrou, visites)",
                "Détenus et visites de son centre",
            ],
            [
                "Médecin",
                "Volet soins (module prévu / en cours)",
                "Consultations médicales",
            ],
        ],
        [4.0, 6.5, 6.5],
    )
    add_caption(doc, "Tableau 2 — Répartition des responsabilités")
    add_body(
        doc,
        "Cette séparation évite qu’un opérateur national écrase, par erreur, les données "
        "d’une prison, tout en lui donnant la visibilité nécessaire au suivi de l’occupation "
        "et des effectifs à l’échelle du pays.",
    )

    # ----- 6 -----
    add_heading_custom(doc, "6.  Chaîne opérationnelle", 1)
    add_body(doc, "Le logiciel s’utilise selon une suite d’étapes qui correspond à la mise en service d’un établissement puis à sa vie quotidienne.")
    add_heading_custom(doc, "6.1  Mise en service d’une prison", 2)
    add_bullet(doc, " l’administrateur central crée le centre (nom, code, type, province, ville, capacité, coordonnées).", "1.")
    add_bullet(doc, " il crée l’administrateur de prison, rattaché à cet établissement, avec identifiant et mot de passe.", "2.")
    add_bullet(doc, " cet administrateur se connecte : le système l’oriente vers le tableau de bord de son centre uniquement.", "3.")
    add_heading_custom(doc, "6.2  Vie de la prison", 2)
    add_bullet(doc, " enregistrement du personnel (directeur, gardiens, infirmiers, agents, etc.).", "4.")
    add_bullet(doc, " écrou du détenu : identité, photos, dossier, tribunal / juridiction, cellule selon le pavillon et le sexe.", "5.")
    add_bullet(doc, " programmation et suivi des visites (visiteur, durée, statut, historique sur la fiche du détenu).", "6.")
    add_bullet(doc, " libération ou transfert, édition de la fiche PDF et de la liste du centre.", "7.")
    add_bullet(doc, " l’administrateur central consulte les mêmes informations à l’échelle nationale, sans les modifier.", "8.")

    # ----- 7 -----
    add_heading_custom(doc, "7.  Fonctionnalités du module", 1)

    add_heading_custom(doc, "7.1  Tableau de bord", 2)
    add_body(
        doc,
        "Le tableau de bord central présente les indicateurs nationaux : nombre de centres, "
        "provinces couvertes, administrateurs, détenus, répartition par statut (incarcéré, "
        "libéré, transféré) et par sexe, capacité et taux d’occupation. Le tableau de bord "
        "d’un centre reprend les mêmes familles d’indicateurs, bornées à l’établissement.",
    )

    add_heading_custom(doc, "7.2  Centres et administrateurs", 2)
    add_body(
        doc,
        "La liste des centres affiche le nom, le code, la ville, la province, le type et "
        "l’occupation. Un filtre par province permet de travailler par entité territoriale. "
        "La fiche d’un centre donne accès aux détenus de cet établissement. Les administrateurs "
        "de prison sont créés et suivis (poste, matricule, affectation, statut actif ou inactif).",
    )

    add_heading_custom(doc, "7.3  Détenus", 2)
    add_body(
        doc,
        "La fiche détenu rassemble l’état civil, les photographies (face et profil), les "
        "informations judiciaires (motif, tribunal choisi dans la liste des juridictions, "
        "numéro et dossier joint), l’affectation pénitentiaire (centre, pavillon, cellule "
        "filtrée selon la structure et le sexe), le régime, les données médicales de base "
        "et l’historique des visites. Les listes sont compactes, avec recherche, pastilles "
        "de statut et actions (consulter, modifier, PDF).",
    )

    add_heading_custom(doc, "7.4  Cellules", 2)
    add_body(
        doc,
        "Les cellules ne sont pas saisies en texte libre. Elles appartiennent à la structure "
        "du centre (pavillons hommes et femmes). Lors de l’écrou, la liste proposée dépend "
        "du centre et du sexe, afin d’éviter une affectation incohérente.",
    )

    add_heading_custom(doc, "7.5  Visites et visiteurs", 2)
    add_body(
        doc,
        "Les visiteurs (famille, avocat, ami, autre) sont enregistrés une fois. Une visite "
        "relie un visiteur à un détenu du centre, avec date, durée et statut (programmée, "
        "en cours, terminée, annulée, refusée). L’historique apparaît sur la fiche du détenu. "
        "Par défaut, la liste des visites met en avant la journée en cours.",
    )

    add_heading_custom(doc, "7.6  Personnel", 2)
    add_body(
        doc,
        "Le module personnel décrit les agents du centre : matricule, fonctions, statut "
        "(actif, congé, inactif). Seuls les responsables du centre créent et mettent à jour "
        "ces fiches.",
    )

    add_heading_custom(doc, "7.7  Surveillance et biométrie", 2)
    add_body(
        doc,
        "Un volet surveillance permet d’associer un profil facial au détenu et d’enregistrer "
        "des détections. La capture biométrique est prévue comme authentification future ; "
        "l’écran correspondant est en place pour les développements ultérieurs. L’identification "
        "principale actuelle reste le compte nominatif et le mot de passe.",
    )

    add_heading_custom(doc, "7.8  Documents officiels", 2)
    add_body(
        doc,
        "La fiche d’identité du détenu et la liste des détenus d’un centre sont éditables "
        "en PDF, avec photographie lorsque celle-ci est disponible. Ces documents sont destinés "
        "aux dossiers d’écrou et aux contrôles administratifs.",
    )

    add_heading_custom(doc, "7.9  Autres domaines du logiciel", 2)
    add_body(
        doc,
        "L’application embarque également des modules de soins, de logistique et de rapports, "
        "accessibles selon le rôle. Ils complètent e-Detenu sans remplacer le circuit judiciaire "
        "des autres fonctions e-Justice, encore à déployer.",
    )

    # ----- 8 -----
    add_heading_custom(doc, "8.  Couverture territoriale", 1)
    add_body(
        doc,
        "Le référentiel initial intègre les prisons centrales et urbaines principales du pays "
        "(trente-cinq établissements), avec province, ville et géolocalisation indicative. "
        "Il ne prétend pas recenser tous les cachots de territoire. L’administrateur central "
        "peut ajouter un centre manquant.",
    )
    add_table(
        doc,
        ["Province / ville", "Établissements (exemples)"],
        [
            ["Kinshasa", "Prison centrale de Makala ; prison militaire de Ndolo"],
            ["Kongo-Central", "Luzumu, Matadi, Boma"],
            ["Kwilu, Kwango", "Kikwit, Bandundu, Kenge"],
            ["Équateur, Tshuapa, Mongala, Ubangi", "Mbandaka, Boende, Angenga, Lisala, Gbadolite, Gemena"],
            ["Tshopo, Uele, Maniema", "Kisangani, Osio, Isiro, Buta, Kindu"],
            ["Nord-Kivu, Sud-Kivu, Ituri", "Munzenze (Goma), Beni, Butembo, Bukavu, Uvira, Bunia"],
            ["Tanganyika, Haut-Katanga, Lualaba, Haut-Lomami", "Kalemie, Kasapa (Lubumbashi), Buluwo (Likasi), Kolwezi, Kamina"],
            ["Kasaï, Lomami, Sankuru", "Kananga, Mbuji-Mayi, Tshikapa, Mwene-Ditu, Lodja"],
        ],
        [7.5, 9.5],
    )
    add_caption(doc, "Tableau 3 — Exemples d’établissements du référentiel national")

    # ----- 9 -----
    add_heading_custom(doc, "9.  Architecture, sécurité et identité visuelle", 1)
    add_heading_custom(doc, "9.1  Technique", 2)
    add_body(
        doc,
        "e-Detenu est une application web (Django) avec une interface de gestion et une API "
        "REST destinée aux intégrations. Les données sont organisées par centre. Les listes, "
        "fiches et formulaires partagent une présentation compacte, adaptée à un usage bureautique "
        "quotidien.",
    )
    add_heading_custom(doc, "9.2  Sécurité d’accès", 2)
    add_bullet(doc, " authentification nominative ; une connexion biométrique est prévue ultérieurement ;", "Compte :")
    add_bullet(doc, " redirection automatique selon le rôle (national ou centre) ;", "Session :")
    add_bullet(doc, " l’administrateur central ne peut pas écrire les données opérationnelles d’une prison ;", "Écriture :")
    add_bullet(doc, " un administrateur de centre ne voit et ne modifie que son établissement.", "Isolation :")
    add_heading_custom(doc, "9.3  Charte graphique", 2)
    add_body(
        doc,
        "L’interface reprend le bleu, le jaune et le rouge du drapeau de la République "
        "démocratique du Congo, avec une présentation institutionnelle (pas de style ludique). "
        "Les statuts (incarcéré, libéré, transféré, etc.) sont distingués par des pastilles lisibles.",
    )

    # ----- 10 -----
    add_heading_custom(doc, "10.  État d’avancement et perspectives", 1)
    add_table(
        doc,
        ["Livré dans e-Detenu", "À développer dans e-Justice"],
        [
            ["Accueil e-Justice et connexion sécurisée", "Police judiciaire et parquet"],
            ["Création des prisons et des admins de centre", "Instruction, greffe, jugements"],
            ["Fiches détenus, dossiers, photos, cellules", "Casier judiciaire national"],
            ["Visites, visiteurs, personnel", "Justice civile et aide juridictionnelle"],
            ["Tableaux de bord national et de centre", "Pilotage statistique transversal"],
            ["PDF (fiche, liste du centre), surveillance (base)", "Biométrie d’authentification aboutie"],
        ],
        [8.5, 8.5],
    )
    add_caption(doc, "Tableau 4 — Ce qui est en service et ce qui reste à bâtir")
    add_body(
        doc,
        "La priorité immédiate reste la consolidation d’e-Detenu dans les centres déjà "
        "dotés d’un administrateur, puis l’extension progressive aux autres établissements "
        "du référentiel, avant d’ouvrir les fonctions judiciaires amont (parquet, greffe).",
    )

    # ----- 11 -----
    add_heading_custom(doc, "11.  Conclusion", 1)
    add_body(
        doc,
        "e-Justice pose le cadre d’une administration judiciaire numérique nationale. "
        "e-Detenu en est le premier maillon vivant : il donne à l’État une vue d’ensemble "
        "des prisons, et à chaque centre un outil de gestion quotidienne, avec une règle "
        "simple — le niveau central crée les établissements et leurs responsables ; "
        "les responsables créent les personnes de leur prison.",
    )
    add_body(
        doc,
        "Cette suite logique — plateforme, module, rôles, chaîne métier, fonctions, territoire, "
        "technique — constitue le socle de présentation du logiciel auprès des autorités "
        "et des équipes de déploiement.",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("—  Fin du document  —")
    set_run_font(r, size=10, italic=True, color=MUTED)

    out = Path(__file__).resolve().parent / "Presentation_e-Justice_e-Detenu.docx"
    doc.save(str(out))
    print(out)


if __name__ == "__main__":
    build()
